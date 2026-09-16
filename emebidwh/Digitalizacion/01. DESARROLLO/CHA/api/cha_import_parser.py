from __future__ import annotations

import hashlib
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class ImportParseError(ValueError):
    pass


def _key(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]+", "", text)


def normalizar_pedido(value: Any) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].replace(".", "", 1).isdigit():
        text = text[:-2]
    return text or None


def normalizar_modelo(texto: Any) -> str:
    return " ".join(str(texto or "").split()).strip()


def clave_comparacion_modelo(texto: Any) -> str:
    value = unicodedata.normalize("NFKC", normalizar_modelo(texto)).casefold()
    return "".join(ch for ch in value if ch.isalnum())


def _is_mark(value: Any) -> bool:
    return isinstance(value, str) and value.strip().casefold() == "x"


def _header(rows, required, aliases=None):
    aliases = aliases or {}
    for row_no, values in rows:
        keys = {_key(v): i for i, v in enumerate(values) if v not in (None, "")}
        found = {}
        for wanted in required:
            candidates = {_key(wanted)} | {_key(x) for x in aliases.get(wanted, set())}
            index = next((i for k, i in keys.items() if k in candidates), None)
            if index is None:
                break
            found[wanted] = index
        if len(found) == len(required):
            return row_no, values, found
    return None


def _rows_for_workbook(stream, filename: str):
    suffix = Path(filename or "").suffix.casefold()
    if suffix not in {".xlsx", ".xls"}:
        raise ImportParseError("Solo se admiten ficheros .xlsx o .xls")
    try:
        if suffix == ".xls":
            import xlrd
            book = xlrd.open_workbook(file_contents=stream.read(), on_demand=True)
            return {sheet.name: [(i + 1, sheet.row_values(i)) for i in range(sheet.nrows)] for sheet in book.sheets()}
        book = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
        return {sheet.title: [(i + 1, list(row)) for i, row in enumerate(sheet.iter_rows(values_only=True))] for sheet in book.worksheets}
    except Exception as exc:
        raise ImportParseError("El fichero no es un Excel legible") from exc


def _metadata(stream, filename: str, sheets, detected: str):
    stream.seek(0)
    content = stream.read()
    return {"nombre": Path(filename or "archivo").name, "tamano": len(content), "sha256": hashlib.sha256(content).hexdigest(), "hojas": list(sheets), "tipo": detected}


def detectar_tipo(sheets) -> str:
    names = {_key(name): name for name in sheets}
    original = all(x in names for x in ("CH1", "CH2", "CEPCC"))
    detalle = any(x in names for x in ("CHH", "CHE", "CC")) and any(_header(rows, ["Ref.Emesa"]) for rows in sheets.values())
    tiempos = "TIEMPOS" in names and bool(_header(sheets[names["TIEMPOS"]], ["CHASIS", "MODELO", "TIEMPO"], {"TIEMPO": {"MINUTOS"}}))
    matches = [kind for kind, ok in (("original_carros", original), ("detalle_cartera", detalle), ("tiempos", tiempos)) if ok]
    if len(matches) != 1:
        raise ImportParseError("No se pudo clasificar inequívocamente el fichero")
    return matches[0]


def parse_file(stream, filename: str):
    if Path(filename or "").suffix.casefold() not in {".xlsx", ".xls"}:
        raise ImportParseError("Extensión no permitida")
    sheets = _rows_for_workbook(stream, filename)
    kind = detectar_tipo(sheets)
    metadata = _metadata(stream, filename, sheets, kind)
    if kind == "original_carros":
        data = _parse_original(sheets)
    elif kind == "detalle_cartera":
        data = _parse_detail(sheets)
    else:
        data = _parse_times(sheets)
    metadata.update(data.pop("metadata", {}))
    return kind, {"metadata": metadata, **data}


def _parse_original(sheets):
    orders = []
    period_weeks, years = set(), set()
    for target, post in (("CH1", "CHM_M1"), ("CH2", "CHM_M2")):
        sheet_name = next((n for n in sheets if _key(n) == target), None)
        header = _header(sheets[sheet_name], ["Nº", "CH.HIDRÁULICO"], {"Nº": {"N", "NUMERO"}, "CH.HIDRÁULICO": {"PEDIDO"}})
        if not header:
            raise ImportParseError(f"Faltan cabeceras obligatorias en {target}")
        row_no, _, indexes = header
        for excel_row, values in sheets[sheet_name][row_no:]:
            pedido = normalizar_pedido(values[indexes["CH.HIDRÁULICO"]] if indexes["CH.HIDRÁULICO"] < len(values) else None)
            visual = normalizar_pedido(values[indexes["Nº"]] if indexes["Nº"] < len(values) else None)
            if not pedido:
                continue
            order = _base_order(values, indexes, excel_row, target, "CHM", visual, pedido)
            order["operaciones"] = [{"fase": "CHM", "puesto": post}]
            orders.append(order)
            _period_values(values, period_weeks, years)
    cep_name = next((n for n in sheets if _key(n) == "CEPCC"), None)
    header = _header(sheets[cep_name], ["Nº", "CH.ELECT", "CH.CONTR"], {"Nº": {"N", "NUMERO"}, "CH.ELECT": {"ELECTRICO"}, "CH.CONTR": {"CONTRAPESO"}})
    if not header:
        raise ImportParseError("Faltan cabeceras obligatorias en CEP,CC")
    row_no, _, indexes = header
    for excel_row, values in sheets[cep_name][row_no:]:
        visual = normalizar_pedido(values[indexes["Nº"]] if indexes["Nº"] < len(values) else None)
        for field, type_code in (("CH.ELECT", "CHE"), ("CH.CONTR", "CHC")):
            pedido = normalizar_pedido(values[indexes[field]] if indexes[field] < len(values) else None)
            if not pedido:
                continue
            order = _base_order(values, indexes, excel_row, "CEP,CC", type_code, visual, pedido)
            order["operaciones"] = [{"fase": "CEP1", "puesto": "CEP1"}, {"fase": "CEP2", "puesto": "CEP2"}]
            orders.append(order)
        _period_values(values, period_weeks, years)
    return {"ordenes": orders, "metadata": {"semanas": sorted(period_weeks), "anios": sorted(years)}}


def _base_order(values, indexes, row, sheet, type_code, visual, pedido):
    def val(name):
        index = indexes.get(name)
        return values[index] if index is not None and index < len(values) else None
    return {"orden_visual": visual, "pedido": pedido, "tipo_chasis": type_code, "fecha_confirmada": _date_value(val("CONFIRMADA")), "montaje": val("MONTAJE"), "embalaje": val("EMBALAJE"), "incidencias": val("INCIDENCIAS"), "hoja_origen": sheet, "fila_origen": row, "errores": [], "avisos": []}


def _date_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip() if value not in (None, "") else None


def _period_values(values, weeks, years):
    for value in values:
        text = str(value or "")
        match = re.search(r"\bS\s*[-_/]?\s*(\d{1,2})\b", text, re.I)
        if match: weeks.add(int(match.group(1)))
        if isinstance(value, (datetime, date)): years.add(value.year)
        else: years.update(int(x) for x in re.findall(r"\b(20\d{2})\b", text))


def _parse_detail(sheets):
    rows_out, weeks, years = [], set(), set()
    relevant = {"CHH": "CHM", "CHE": "CHE", "CC": "CHC"}
    found_sheet = False
    for sheet_name, type_code in relevant.items():
        actual = next((n for n in sheets if _key(n) == sheet_name), None)
        if not actual: continue
        header = _header(sheets[actual], ["Ref.Emesa"], {"Ref.Emesa": {"REFEMESA", "REF"}})
        if not header: continue
        found_sheet = True
        row_no, header_values, indexes = header
        ref_index = indexes["Ref.Emesa"]
        ignored = {"REFEMESA", "REF", "CLIENTE", "CLIENTEEXTERNO", "PEDIDO"}
        model_indexes = [(i, normalizar_modelo(v)) for i, v in enumerate(header_values) if normalizar_modelo(v) and _key(v) not in ignored and i != ref_index]
        for excel_row, values in sheets[actual][row_no:]:
            pedido = normalizar_pedido(values[ref_index] if ref_index < len(values) else None)
            if not pedido: continue
            marked = [model for i, model in model_indexes if i < len(values) and _is_mark(values[i])]
            client_index = next((i for i, v in enumerate(header_values) if _key(v) in {"CLIENTE", "CLIENTEEXTERNO"}), None)
            rows_out.append({"pedido": pedido, "tipo_chasis": type_code, "modelo": marked[0] if len(marked) == 1 else None, "marcas": marked, "cliente": values[client_index] if client_index is not None and client_index < len(values) else None, "hoja_origen": actual, "fila_origen": excel_row, "errores": (["Ninguna marca de modelo"] if not marked else ["Más de una marca de modelo"] if len(marked) > 1 else [])})
            _period_values(values, weeks, years)
    if not found_sheet: raise ImportParseError("Faltan hojas/cabeceras de Detalle Cartera")
    return {"detalle": rows_out, "metadata": {"semanas": sorted(weeks), "anios": sorted(years)}}


def _parse_times(sheets):
    actual = next((n for n in sheets if _key(n) == "TIEMPOS"), None)
    header = _header(sheets[actual], ["CHASIS", "MODELO", "TIEMPO"], {"TIEMPO": {"MINUTOS"}})
    if not header: raise ImportParseError("Faltan cabeceras obligatorias en Tiempos")
    row_no, header_values, indexes = header
    all_indexes = {_key(v): i for i, v in enumerate(header_values)}
    phase_indexes = {"CEP1": all_indexes.get("CEP1"), "CEP2": all_indexes.get("CEP2")}
    result, seen = [], set()
    for excel_row, values in sheets[actual][row_no:]:
        chassis = normalizar_modelo(values[indexes["CHASIS"]] if indexes["CHASIS"] < len(values) else "").upper()
        model = normalizar_modelo(values[indexes["MODELO"]] if indexes["MODELO"] < len(values) else "")
        if not chassis and not model: continue
        type_code = next((x for x in ("CHM", "CHE", "CHC") if x in chassis), chassis)
        if type_code == "CHM": phase_values = {"CHM": values[indexes["TIEMPO"]] if indexes["TIEMPO"] < len(values) else None}
        else: phase_values = {phase: values[index] if index is not None and index < len(values) else None for phase, index in phase_indexes.items()}
        if not model: phase_values = {None: None}
        for phase, raw in phase_values.items():
            try: minutes = float(raw)
            except (TypeError, ValueError): minutes = None
            error = None if minutes is not None and minutes > 0 else "El tiempo debe ser numérico y mayor que 0"
            identity = (type_code, clave_comparacion_modelo(model), phase)
            if identity in seen: error = "Duplicado incompatible en Tiempos"
            seen.add(identity)
            result.append({"tipo_chasis": type_code, "modelo": model, "fase": phase, "minutos": minutes, "fila_origen": excel_row, "error": error})
    return {"tiempos": result}
